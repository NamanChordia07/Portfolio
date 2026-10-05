"""Build the resumes from content.yaml: three role-specific versions for applications, and one combined
version for the website.

    python build.py            # all variants: dist/*.pdf, dist/*.docx, dist/*.html, ../public/resume/*.pdf
    python build.py --check    # also fail if any PDF is longer than one page

Outputs
-------
dist/<file>.pdf    application copy
dist/<file>.docx   editable Word copy (for portals that want .docx)
dist/<file>.html   the exact HTML the PDF is printed from
../public/resume/<file>.pdf   website copy (variants marked ``public: true``)
"""

from __future__ import annotations

import argparse
import html
import os
import shutil
import sys
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).parent
DIST = ROOT / "dist"
PUBLIC = ROOT.parent / "public" / "resume"
CHROME = os.environ.get("CHROME_PATH", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome")


def load() -> dict[str, Any]:
    return yaml.safe_load((ROOT / "content.yaml").read_text(encoding="utf-8"))


def pick(item: dict[str, Any], variant: str, fallback: str | None = None) -> str | None:
    """The bullet text for a variant, or None if the bullet is excluded.

    A variant with a ``fallback`` (the combined resume) borrows that variant's wording and inclusion where it has
    none of its own.
    """
    if "in" in item and variant not in item["in"] and (fallback is None or fallback not in item["in"]):
        return None
    text = item.get(variant) or (item.get(fallback) if fallback else None) or item.get("text")
    return " ".join(text.split()) if text else None


def ranked(items: list[dict[str, Any]], variant: str, fallback: str | None = None) -> list[dict[str, Any]]:
    def key(i: dict[str, Any]) -> int:
        rank = i.get("rank", {})
        return int(rank.get(variant, rank.get(fallback, 99) if fallback else 99))

    return sorted(items, key=key)


# --- structure shared by the HTML and DOCX renderers ----------------------------------------


def assemble(data: dict[str, Any], variant: str) -> dict[str, Any]:
    v = data["variants"][variant]
    fb = v.get("fallback")
    person = data["person"]
    line1 = [person["location"], person["phone"], person["email"]]
    contact = [line1, [link["label"] for link in person["links"]]]
    experience = []
    for org in data["experience"]:
        roles = []
        short = org["company"].split(" (")[0]
        for i, role in enumerate(org["roles"]):
            bullets = [t for b in ranked(role["bullets"], variant, fb) if (t := pick(b, variant, fb))]
            # later roles at the same company name it, so parsers do not merge them into the first role
            display = role["title"] if i == 0 else f"{role['title']}, {short}"
            roles.append({"title": role["title"], "display": display, "dates": role["dates"], "bullets": bullets})
        experience.append({"company": org["company"], "location": org["location"], "roles": roles})
    projects = []
    for p in ranked(data["projects"], variant, fb):
        bullets = [t for b in p["bullets"] if (t := pick(b, variant, fb))]
        projects.append({**p, "bullets": bullets})
    return {
        "name": person["name"],
        "headline": v["headline"],
        "contact": contact,
        "links": {link["label"]: link["url"] for link in person["links"]} | {"email": f"mailto:{person['email']}"},
        "summary": " ".join(v["summary"].split()),
        "sections": v["sections"],
        "experience": experience,
        "projects": projects,
        "skills": data["skills"].get(variant) or data["skills"][fb],
        "education": data["education"],
        "leadership": [item["text"] for item in data["leadership"]],
    }


# --- HTML (printed to PDF by Chromium) --------------------------------------------------------

CSS = """
@font-face { font-family: "Source Sans 3"; font-weight: 400; src: url("fonts/source-sans-3-latin-400-normal.woff2") format("woff2"); }
@font-face { font-family: "Source Sans 3"; font-weight: 400; font-style: italic; src: url("fonts/source-sans-3-latin-400-italic.woff2") format("woff2"); }
@font-face { font-family: "Source Sans 3"; font-weight: 600; src: url("fonts/source-sans-3-latin-600-normal.woff2") format("woff2"); }
@font-face { font-family: "Source Sans 3"; font-weight: 700; src: url("fonts/source-sans-3-latin-700-normal.woff2") format("woff2"); }
@page { size: A4; margin: 0.45in 0.55in 0.4in; }
* { box-sizing: border-box; }
html { font-family: "Source Sans 3", Arial, sans-serif; font-size: 10.1pt; line-height: 1.27; color: #000; background: #fff; }
body { margin: 0; }
header { text-align: center; }
h1 { font-size: 21pt; font-weight: 700; margin: 0; letter-spacing: 0.4pt; line-height: 1.1; }
.headline { font-size: 10.8pt; font-weight: 600; margin-top: 1.5pt; }
.contact { font-size: 9.6pt; margin-top: 1.5pt; }
.contact span + span::before { content: " | "; }
h2 { font-size: 10.6pt; font-weight: 700; text-transform: uppercase; letter-spacing: 0.9pt;
     border-bottom: 0.7pt solid #000; margin: 7.5pt 0 3pt; padding-bottom: 0.8pt; }
.summary { margin: 0; }
.row { display: flex; justify-content: space-between; align-items: baseline; gap: 12pt; }
.row .right { white-space: nowrap; }
.org { font-weight: 700; }
.role { font-weight: 600; margin-top: 1.5pt; }
.entry { margin-bottom: 3pt; }
.entry + .entry { margin-top: 2pt; }
.tech { font-style: italic; }
ul { margin: 1pt 0 1pt; padding-left: 13pt; }
li { margin: 0 0 1.3pt; padding-left: 1pt; }
.skills div { margin-bottom: 1pt; }
.skills b { font-weight: 600; }
a { color: inherit; text-decoration: none; }
"""


def e(s: str) -> str:
    return html.escape(s, quote=True)


def link(label: str, links: dict[str, str]) -> str:
    url = links.get(label)
    return f'<a href="{e(url)}">{e(label)}</a>' if url else e(label)


def render_html(r: dict[str, Any]) -> str:
    L = r["links"]
    def contact_item(c: str) -> str:
        if c in L:
            return link(c, L)
        if "@" in c:
            return f'<a href="{e(L["email"])}">{e(c)}</a>'
        return e(c)

    contact = "".join(
        "<div class='contact'>" + "".join(f"<span>{contact_item(c)}</span>" for c in line) + "</div>" for line in r["contact"]
    )
    parts = [
        f"<header><h1>{e(r['name'])}</h1><div class='headline'>{e(r['headline'])}</div>"
        f"{contact}</header>"
    ]
    for section in r["sections"]:
        if section == "summary":
            parts.append(f"<section><h2>Summary</h2><p class='summary'>{e(r['summary'])}</p></section>")
        elif section == "experience":
            items = []
            for org in r["experience"]:
                roles = []
                for i, role in enumerate(org["roles"]):
                    head = (
                        f"<div class='row'><span class='org'>{e(org['company'])}</span><span class='right'>{e(org['location'])}</span></div>"
                        if i == 0
                        else ""
                    )
                    bullets = "".join(f"<li>{e(b)}</li>" for b in role["bullets"])
                    roles.append(
                        f"{head}<div class='row role'><span>{e(role['display'])}</span><span class='right'>{e(role['dates'])}</span></div>"
                        f"<ul>{bullets}</ul>"
                    )
                items.append(f"<div class='entry'>{''.join(roles)}</div>")
            parts.append(f"<section><h2>Experience</h2>{''.join(items)}</section>")
        elif section == "projects":
            items = []
            for p in r["projects"]:
                right = f"{link(p['link']['label'], {p['link']['label']: p['link']['url']})} | {e(p['dates'])}"
                bullets = "".join(f"<li>{e(b)}</li>" for b in p["bullets"])
                items.append(
                    f"<div class='entry'><div class='row'><span><span class='org'>{e(p['name'])}</span> | "
                    f"<span class='tech'>{e(p['tech'])}</span></span><span class='right'>{right}</span></div>"
                    f"<ul>{bullets}</ul></div>"
                )
            parts.append(f"<section><h2>Projects</h2>{''.join(items)}</section>")
        elif section == "skills":
            rows = "".join(f"<div><b>{e(k)}:</b> {e(v)}</div>" for k, v in r["skills"])
            parts.append(f"<section class='skills'><h2>Technical Skills</h2>{rows}</section>")
        elif section == "education":
            rows = "".join(
                f"<div class='row'><span><span class='org'>{e(ed['school'])}</span> | {e(ed['degree'])}</span>"
                f"<span class='right'>{e(ed['dates'])}</span></div>"
                for ed in r["education"]
            )
            parts.append(f"<section><h2>Education</h2>{rows}</section>")
        elif section == "leadership":
            bullets = "".join(f"<li>{e(t)}</li>" for t in r["leadership"])
            parts.append(f"<section><h2>Leadership &amp; Awards</h2><ul>{bullets}</ul></section>")
    title = f"{r['name']} - Resume"
    return (
        f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(title)}</title>"
        f"<meta name='author' content='{e(r['name'])}'><style>{CSS}</style></head><body>{''.join(parts)}</body></html>"
    )


def print_pdfs(jobs: list[tuple[Path, Path]]) -> None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROME)
        page = browser.new_page()
        for html_path, pdf_path in jobs:
            page.goto(html_path.resolve().as_uri())
            page.evaluate("document.fonts.ready")
            page.pdf(path=str(pdf_path), prefer_css_page_size=True, print_background=False, tagged=True, outline=False)
        browser.close()


# --- DOCX -------------------------------------------------------------------------------------


def render_docx(r: dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
    sec.left_margin = sec.right_margin = Inches(0.55)
    sec.top_margin, sec.bottom_margin = Inches(0.45), Inches(0.4)
    text_width = sec.page_width - sec.left_margin - sec.right_margin
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def para(text: str = "", bold: bool = False, size: float | None = None, align=None, space_before: float = 0):  # type: ignore[no-untyped-def]
        p = doc.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        return p

    def row(left: str, right: str, bold_left: bool = True, italic_tail: str = "") -> None:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run(left)
        run.bold = bold_left
        if italic_tail:
            t = p.add_run(" | " + italic_tail)
            t.italic = True
        p.add_run("\t" + right)

    def heading(text: str) -> None:
        p = para(text.upper(), bold=True, size=10.5, space_before=6)
        p.paragraph_format.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(k), v)
        border.append(bottom)
        pPr.append(border)

    def bullets(items: list[str]) -> None:
        for b in items:
            p = doc.add_paragraph(b, style="List Bullet")
            p.paragraph_format.space_after = Pt(1)

    para(r["name"], bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(r["headline"], bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in r["contact"]:
        para(" | ".join(line), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for section in r["sections"]:
        if section == "summary":
            heading("Summary")
            para(r["summary"])
        elif section == "experience":
            heading("Experience")
            for org in r["experience"]:
                row(org["company"], org["location"])
                for role in org["roles"]:
                    row(role["display"], role["dates"], bold_left=False)
                    bullets(role["bullets"])
        elif section == "projects":
            heading("Projects")
            for p in r["projects"]:
                row(p["name"], f"{p['link']['label']} | {p['dates']}", italic_tail=p["tech"])
                bullets(p["bullets"])
        elif section == "skills":
            heading("Technical Skills")
            for k, v in r["skills"]:
                p = doc.add_paragraph()
                p.add_run(f"{k}: ").bold = True
                p.add_run(v)
        elif section == "education":
            heading("Education")
            for ed in r["education"]:
                row(ed["school"], ed["dates"], italic_tail="")
                para(ed["degree"])
        elif section == "leadership":
            heading("Leadership & Awards")
            bullets(r["leadership"])
    doc.core_properties.author = r["name"]
    doc.core_properties.title = f"{r['name']} - Resume"
    doc.save(str(path))


# --- main -------------------------------------------------------------------------------------


def previews(pdfs: list[Path]) -> None:
    """PNG thumbnails of the public PDFs for the website's resume page (794 px wide, A4 at 96 dpi)."""
    import pypdfium2 as pdfium

    out = PUBLIC / "previews"
    out.mkdir(parents=True, exist_ok=True)
    for pdf in pdfs:
        page = pdfium.PdfDocument(str(pdf))[0]
        image = page.render(scale=96 / 72).to_pil().convert("L")
        image.save(out / f"{pdf.stem}.png", optimize=True)


def page_count(pdf: Path) -> int:
    from pypdf import PdfReader

    return len(PdfReader(str(pdf)).pages)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="fail if any PDF exceeds one page")
    ap.add_argument("--variants", default="fde,ai,sde,web")
    args = ap.parse_args()
    data = load()
    DIST.mkdir(exist_ok=True)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    jobs = []
    public = []
    for variant in args.variants.split(","):
        name = data["variants"][variant]["file"]
        r = assemble(data, variant)
        html_path = DIST / f"{name}.html"
        html_path.write_text(render_html(r).replace('url("fonts/', 'url("../fonts/'), encoding="utf-8")
        jobs.append((html_path, DIST / f"{name}.pdf"))
        render_docx(r, DIST / f"{name}.docx")
        if data["variants"][variant].get("public"):
            public.append(PUBLIC / f"{name}.pdf")
    print_pdfs(jobs)
    for pdf in public:
        shutil.copyfile(DIST / pdf.name, pdf)
    previews(public)
    failed = False
    for pdf in [pdf for _, pdf in jobs] + public:
        n = page_count(pdf)
        print(f"{pdf.relative_to(ROOT.parent)}: {n} page(s)")
        failed |= n != 1
    return 1 if (failed and args.check) else 0


if __name__ == "__main__":
    sys.exit(main())
